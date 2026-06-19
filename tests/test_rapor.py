# Rapor üretimi: formül doğruluğu, DIV/0 koruması, ürün kolonları, şablon.
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import pytest

from lazer import rapor


def _satir(magaza, ciro, vergi, reklam, km, kargo, adet, ilave=0, upgrade=0,
           urunler=None):
    return {"magaza": magaza, "ciro": ciro, "vergi": vergi, "reklam": reklam,
            "kargo_musteri": km, "kargo": kargo, "adet": adet,
            "ilave_odeme": ilave, "upgrade": upgrade, "urunler": urunler or {}}


@pytest.fixture
def basit_rapor(tmp_path):
    yol = str(tmp_path / "r.xlsx")
    satirlar = [
        _satir("Mağaza A", 1000, 100, 50, 200, 150, 100, ilave=10, upgrade=5,
               urunler={"Coffee Mug": 12, "Wallet": 3}),
        _satir("Mağaza B", 5000, 500, 300, 800, 600, 250, urunler={"Coffee Mug": 40}),
    ]
    rapor.rapor_uret(yol, "Mayıs", 2026, satirlar,
                     urun_basliklari=["Coffee Mug", "Wallet"])
    return yol


def _hesapla(yol):
    import formulas
    sol = formulas.ExcelModel().loads(yol).finish().calculate()
    k = {}
    for key, v in sol.items():
        m = re.search(r"RAPOR'!([A-Z]+\d+)\b", key)
        if not m:
            continue
        try:
            k[m.group(1)] = v.value[0, 0]
        except Exception:
            pass
    return k


def test_baslik_ve_sablon(basit_rapor):
    ws = openpyxl.load_workbook(basit_rapor).active
    assert ws.cell(1, 1).value == "MAĞAZA"
    assert ws.cell(1, 1).fill.fgColor.rgb.endswith("FF9900")   # turuncu
    assert ws.cell(1, 1).font.bold is True


def test_finansal_kolon_sirasi(basit_rapor):
    ws = openpyxl.load_workbook(basit_rapor).active
    bekl = ["MAĞAZA", "CİRO", "Parça başı ciro", "VERGİ", "REKLAM",
            "KARGO MÜŞTERİ"]
    for i, ad in enumerate(bekl, start=1):
        assert ws.cell(1, i).value == ad


def test_formuller_hatasiz(basit_rapor):
    k = _hesapla(basit_rapor)
    hata = [(c, v) for c, v in k.items()
            if isinstance(v, str) and v.startswith("#")]
    assert hata == []


def test_kalan_formulu(basit_rapor):
    k = _hesapla(basit_rapor)
    # KALAN = CİRO - VERGİ - REKLAM - Kargo (A satırı = 3. satır)
    assert abs(float(k["M3"]) - (1000 - 100 - 50 - 150)) < 0.01


def test_parca_basi_ciro(basit_rapor):
    k = _hesapla(basit_rapor)
    assert abs(float(k["C3"]) - (1000 / 100)) < 0.01


def test_kargo_pb_kalan_formulu(basit_rapor):
    k = _hesapla(basit_rapor)
    # (KARGO MÜŞTERİ - Kargo) / ADET = (200-150)/100
    assert abs(float(k["H3"]) - ((200 - 150) / 100)) < 0.01


def test_yuzde_kolonlari(basit_rapor):
    k = _hesapla(basit_rapor)
    assert abs(float(k["N3"]) - (50 / 1000)) < 0.0001   # %REKLAM
    assert abs(float(k["O3"]) - (100 / 1000)) < 0.0001  # %VERGİ


def test_toplam_satiri(basit_rapor):
    k = _hesapla(basit_rapor)
    # TOPLAM CİRO = 1000 + 5000
    toplam_satir = None
    for r in range(3, 12):
        if str(k.get(f"A{r}", "")) == "TOPLAM":
            toplam_satir = r
            break
    assert toplam_satir is not None
    assert abs(float(k[f"B{toplam_satir}"]) - 6000) < 0.01


def test_div0_korumasi(tmp_path):
    yol = str(tmp_path / "z.xlsx")
    rapor.rapor_uret(yol, "Mayıs", 2026, [
        _satir("Sıfır Adet", 100, 10, 5, 0, 0, 0),   # adet=0
        _satir("Sıfır Ciro", 0, 0, 0, 0, 0, 0),      # ciro=0
    ])
    k = _hesapla(yol)
    hata = [(c, v) for c, v in k.items()
            if isinstance(v, str) and v.startswith("#")]
    assert hata == []


def test_urun_kolonlari(basit_rapor):
    ws = openpyxl.load_workbook(basit_rapor).active
    # 17 finansal + 2 ürün = 19 kolon
    assert ws.cell(1, 18).value == "Coffee Mug"
    assert ws.cell(1, 19).value == "Wallet"
    assert ws.cell(3, 18).value == 12   # Mağaza A coffee mug
    assert ws.cell(4, 18).value == 40   # Mağaza B coffee mug


def test_urun_toplami(basit_rapor):
    k = _hesapla(basit_rapor)
    for r in range(3, 12):
        if str(k.get(f"A{r}", "")) == "TOPLAM":
            assert abs(float(k[f"R{r}"]) - (12 + 40)) < 0.01   # Coffee Mug toplam
            break


def test_amazon_bolumu(tmp_path):
    yol = str(tmp_path / "a.xlsx")
    rapor.rapor_uret(yol, "Mayıs", 2026,
                     [_satir("A", 100, 10, 5, 20, 15, 10)],
                     amazon_satirlari=[{"magaza": "Foma Amazon Store",
                                        "siparis": 100, "kargo": 500.0}])
    ws = openpyxl.load_workbook(yol).active
    bulundu = any(ws.cell(r, 1).value == "RAPOR DIŞI (AMAZON)"
                  for r in range(1, ws.max_row + 1))
    assert bulundu


def test_para_format(basit_rapor):
    ws = openpyxl.load_workbook(basit_rapor).active
    assert ws.cell(3, 2).number_format == "#,##0"     # CİRO
    assert ws.cell(3, 14).number_format == "0.0%"     # %REKLAM
