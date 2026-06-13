# Yeni ShipStation Sipariş Export'u (maliyet dahil) parser + uçtan uca testleri.
# Onaylı dedupe (gelir drop_duplicates, kargo distinct-per-order), pazar yeri
# ayrımı, iptal/sıfır şüphesi, boş Order #, 0-item, ad normalizasyonu, en-erken
# tarih ay ataması ve Excel çıktısı doğrulanır.
import csv
import os
import shutil
import sys
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from lazer import shipstation_csv

H = ["Order - Number", "Date - Order Date", "Market - Store Name",
     "Market - Market Order URL", "Market - Markeplace Name",
     "Amount - Shipping Cost", "Amount - Order Shipping",
     "Amount - Paid by Customer", "Amount - Order Subtotal",
     "Amount - Order Tax", "Amount - Order Total", "Count - Number of Items",
     "Custom - Field 1"]

V = "Velvet Whiskey Design"
VPAD = "  Velvet  Whiskey   Design "   # baş/son + iç tekrarlı boşluk
CC = "Crafty Corner Shop"
PO = "Pine & Oak Studio"


def _row(no, date, store, mp, ship, oship, paid, sub, tax, total, items):
    url = "https://etsy.com/o/" + str(no) if mp.lower() == "etsy" else ""
    return [no, date, store, url, mp, ship, oship, paid, sub, tax, total, items, ""]


def write_export(path):
    rows = [
        # Velvet — tek sipariş, AM/PM tarih, padded ad
        _row("E-1", "05/01/2026 10:30:00 AM", VPAD, "Etsy", 8.50, 10, 250, 230, 20, 250, 3),
        # Velvet — GERÇEK kopya: aynı kargo 6.00, farklı tarihler -> kargo 1 kez
        _row("E-2", "05/05/2026 09:00:00 AM", V, "Etsy", 6.00, 5, 100, 95, 8, 100, 2),
        _row("E-2", "05/20/2026 09:00:00 AM", V, "Etsy", 6.00, 5, 100, 95, 8, 100, 2),
        # Crafty — ÇOK GÖNDERİLİ: farklı kargolar 4.00 + 3.00 -> ikisi de sayılır
        _row("E-3", "05/10/2026 02:15:00 PM", CC, "Etsy", 4.00, 8, 200, 192, 15, 200, 4),
        _row("E-3", "05/10/2026 02:15:00 PM", CC, "Etsy", 3.00, 8, 200, 192, 15, 200, 4),
        # Crafty — son gün
        _row("E-4", "05/31/2026", CC, "Etsy", 5.25, 5, 90, 83, 7, 90, 1),
        # Velvet — NEGATİF (iptal şüphesi)
        _row("E-5", "05/15/2026", V, "Etsy", 0, 0, -30, -30, 0, -30, 1),
        # Crafty — SIFIR total (iptal şüphesi)
        _row("E-6", "05/16/2026", CC, "Etsy", 0, 0, 0, 0, 0, 0, 1),
        # Pine & Oak — 0 item
        _row("E-7", "05/12/2026", PO, "Etsy", 4.75, 3, 60, 55, 5, 60, 0),
        # Lonely — formu yok (satış var form yok listesine düşer)
        _row("E-8", "05/08/2026", "Lonely Etsy Store", "Etsy", 9.00, 7, 150, 138, 12, 150, 2),
        # Crafty — BOŞ Order # (dedupe'a girmez, ayrı sipariş)
        _row("", "05/09/2026", CC, "Etsy", 2.00, 2, 45, 43, 0, 45, 1),
        # Amazon — pazar yeri authoritative (ad'da amazon var)
        _row("111-2233445-6677889", "05/14/2026", "Amazon US - LazerCo", "Amazon", 12.00, 0, 75, 75, 0, 75, 1),
        # Amazon — ad'da amazon YOK ama marketplace Amazon -> yine Amazon
        _row("111-9988776-5544332", "05/18/2026", "Some Store", "Amazon", 3.00, 0, 40, 40, 0, 40, 1),
        # Fallback — marketplace boş, ad'da amazon -> isim sezgisiyle Amazon
        _row("X-1", "05/19/2026", "Amazon Handmade Co", "", 4.00, 0, 50, 50, 0, 50, 1),
        # Ay dışı: Nisan ve Haziran
        _row("E-9", "04/30/2026", CC, "Etsy", 5.00, 5, 70, 65, 5, 70, 1),
        _row("E-10", "06/01/2026", V, "Etsy", 5.00, 5, 55, 50, 4, 55, 1),
        # EN ERKEN tarih ay ataması: kopya satırlar 04/29 + 05/03 -> Nisan, hariç
        _row("E-11", "04/29/2026", V, "Etsy", 5.00, 5, 100, 95, 8, 100, 1),
        _row("E-11", "05/03/2026", V, "Etsy", 5.00, 5, 100, 95, 8, 100, 1),
    ]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(H)
        w.writerows(rows)


class TestParser(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.csv = os.path.join(cls.tmp, "shipstation_orders.csv")
        write_export(cls.csv)
        cls.r = shipstation_csv.ozet_isle(cls.csv, 5, 2026)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_store_name_normalized(self):
        # Padded ' Velvet  Whiskey   Design ' tek anahtarda toplanır
        self.assertIn(V, self.r["magazalar"])
        self.assertNotIn(VPAD.strip(), [k for k in self.r["magazalar"] if k != V])

    def test_revenue_deduped_per_order(self):
        m = self.r["magazalar"]
        # Velvet: E-1 + E-2 (kopya, 1 kez). E-5 neg, E-10/E-11 ay dışı hariç.
        self.assertAlmostEqual(m[V]["ciro_order_total"], 350.0)
        self.assertAlmostEqual(m[V]["vergi"], 28.0)
        self.assertAlmostEqual(m[V]["kargo_musteri"], 15.0)
        self.assertEqual(m[V]["adet"], 5.0)
        self.assertEqual(m[V]["siparis_sayisi"], 2)

    def test_kargo_distinct_vs_naive(self):
        m = self.r["magazalar"]
        # Velvet kargo: distinct 8.50 + 6.00 = 14.50 (kopya 6.00 çift sayılmaz)
        self.assertAlmostEqual(m[V]["kargo_maliyet"], 14.50)
        # Crafty kargo: E-3 (4+3=7) + E-4 5.25 + boş 2 = 14.25
        self.assertAlmostEqual(m[CC]["kargo_maliyet"], 14.25)
        # Dedupe kanıtı: naive, gerçek kopya yüzünden distinct'ten 6.00 fazla
        self.assertAlmostEqual(self.r["kargo_naive_toplam"] - self.r["kargo_distinct_toplam"], 6.00)

    def test_marketplace_split_authoritative(self):
        m, amz = self.r["magazalar"], self.r["amazon"]
        # Hiçbir Amazon mağazası gövdeye girmedi
        for amazon_store in ("Amazon US - LazerCo", "Some Store", "Amazon Handmade Co"):
            self.assertNotIn(amazon_store, m)
        self.assertEqual(self.r["pazar_sayilari"]["amazon"], 3)
        # 'Some Store' adı amazon içermez ama marketplace Amazon -> Amazon
        self.assertIn("Some Store", amz)
        self.assertAlmostEqual(amz["Amazon US - LazerCo"]["kargo"], 12.0)
        self.assertEqual(sum(v["siparis"] for v in amz.values()), 3)

    def test_refund_suspects_excluded(self):
        # E-5 negatif, E-6 sıfır -> gelirden hariç, per-store sayım
        suspects = {x["order_no"] for x in self.r["iptal_iade"]}
        self.assertIn("E-5", suspects)
        self.assertIn("E-6", suspects)
        self.assertEqual(self.r["iptal_magaza"].get(V), 1)
        self.assertEqual(self.r["iptal_magaza"].get(CC), 1)

    def test_blank_order_no_separate(self):
        b = self.r["bos_order_no"][CC]
        self.assertEqual(b["satir"], 1)
        self.assertAlmostEqual(b["ciro"], 45.0)
        # Boş sipariş Crafty cirosuna dahil: 200 + 90 + 45 = 335
        self.assertAlmostEqual(self.r["magazalar"][CC]["ciro_order_total"], 335.0)

    def test_zero_item_kept(self):
        self.assertEqual(self.r["magazalar"][PO]["adet"], 0.0)
        self.assertAlmostEqual(self.r["magazalar"][PO]["ciro_order_total"], 60.0)

    def test_earliest_date_month_assignment(self):
        # E-11 en erken 04/29 -> Nisan -> Mayıs'a girmemeli (Velvet ciro 350, 450 değil)
        self.assertAlmostEqual(self.r["magazalar"][V]["ciro_order_total"], 350.0)
        self.assertNotIn("E-11", self.r["siparisler"])

    def test_ampm_date_parsed(self):
        t = shipstation_csv._tarih("05/01/2026 10:30:00 AM")
        self.assertEqual((t.month, t.day), (5, 1))
        t2 = shipstation_csv._tarih("05/12/2026 02:15:00 PM")
        self.assertEqual((t2.month, t2.day, t2.hour), (5, 12, 14))


class TestEndToEnd(unittest.TestCase):
    def setUp(self):
        self.eski = os.getcwd()
        self.tmp = tempfile.mkdtemp()
        os.chdir(self.tmp)
        write_export(os.path.join(self.tmp, "export.csv"))
        # master form (3 eşleşen mağaza)
        master = pd.DataFrame([
            dict(zip(["Mağaza","Dönem","Yıl","Ciro","Vergi","Reklam",
                      "Müşteri Kargo Ödemesi","Kargo Harcaması","Toplam Satış Adedi",
                      "Upgrade Satışı","Ek ödeme"], r))
            for r in [
                (V, "Mayıs", 2026, 360, 30, 25, 16, 15, 6, 1, "12$ Adobe"),
                ("Crafty Corner", "Mayıs", 2026, 320, 23, 15, 14, 14, 6, 0, ""),
                ("Pine & Oak Studio", "Mayıs", 2026, 60, 5, 0, 3, 5, 1, 0, ""),
            ]])
        with pd.ExcelWriter("master.xlsx") as w:
            master.to_excel(w, sheet_name="Form Yanıtları 1", index=False)

    def tearDown(self):
        os.chdir(self.eski)
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_full_flow_excel(self):
        import importlib
        import app as APP
        importlib.reload(APP)  # taze DURUM + cwd-göreli yollar
        from lazer import eslestirme
        eslestirme.kaydet({V: V, CC: "Crafty Corner", PO: "Pine & Oak Studio"})
        APP.DURUM["ozet_yolu"] = os.path.join(self.tmp, "export.csv")
        c = APP.app.test_client()
        j = c.post("/api/analiz", json={"ay": "Mayıs", "yil": 2026}).get_json()
        self.assertTrue(j.get("ozet_var"))
        self.assertEqual(j["form_magaza_sayisi"], 3)
        # P kolonu + ürün adedi + dedupe notları analiz uyarılarında
        self.assertTrue(any("P kolonu boş bırakıldı" in u for u in j["uyarilar"]))
        self.assertTrue(any("SKU/ürün kalemi yok" in u for u in j["uyarilar"]))
        self.assertTrue(any("Dedupe kanıtı" in u for u in j["uyarilar"]))

        d = c.post("/api/denetim", json={}).get_json()
        # Lonely Etsy Store: satış var form yok, Ciro+Kargo ile
        lonely = [x for x in d["eslesmeyen_detay"] if x["magaza"] == "Lonely Etsy Store"]
        self.assertTrue(lonely and abs(lonely[0]["ciro"] - 150.0) < 0.005
                        and abs(lonely[0]["kargo"] - 9.0) < 0.005)
        # Amazon rapor dışı
        self.assertTrue(any("Amazon" in a["magaza"] for a in d["amazon"]))

        r = c.post("/api/rapor", json={"kaynaklar": {}}).get_json()
        self.assertTrue(r.get("tamam"))
        wb = load_workbook(APP.DURUM["son_rapor"]); ws = wb.active
        rows = {}
        rr = 3
        while ws.cell(row=rr, column=1).value not in (None, "TOPLAM"):
            rows[ws.cell(row=rr, column=1).value] = rr; rr += 1
        # Header P kolonu korunmuş + boş
        self.assertEqual(ws.cell(row=1, column=10).value, "P")
        self.assertTrue(all(ws.cell(row=x, column=10).value is None for x in rows.values()))
        # ShipStation-bağlı değerler (kaynaklar={} -> hepsi shipstation)
        exp = {  # ciro(B2), vergi(D4), km(F6), kargo(I9), adet(K11)
            "Crafty Corner": (335, 22, 15, 14.25, 6),
            "Pine & Oak Studio": (60, 5, 3, 4.75, 0),
            V: (350, 28, 15, 14.50, 5),
        }
        for st, (ci, ve, km, ka, ad) in exp.items():
            x = rows[st]
            self.assertAlmostEqual(ws.cell(row=x, column=2).value, ci, msg=st+" ciro")
            self.assertAlmostEqual(ws.cell(row=x, column=4).value, ve, msg=st+" vergi")
            self.assertAlmostEqual(ws.cell(row=x, column=6).value, km, msg=st+" km")
            self.assertAlmostEqual(ws.cell(row=x, column=9).value, ka, msg=st+" kargo")
            self.assertEqual(ws.cell(row=x, column=11).value, ad, msg=st+" adet")
            # KALAN formülü (gerçek Excel formülü)
            self.assertEqual(ws.cell(row=x, column=13).value, f"=B{x}-D{x}-E{x}-I{x}")
        # Reklam/İlave/Upgrade formdan
        vrow = rows[V]
        self.assertAlmostEqual(ws.cell(row=vrow, column=5).value, 25)   # reklam
        self.assertAlmostEqual(ws.cell(row=vrow, column=12).value, 12)  # ilave (12$ Adobe)
        self.assertEqual(ws.cell(row=vrow, column=17).value, 1)         # upgrade
        # 0-item Pine: parça başı ciro hücresi DIV/0 guard'lı formül
        prow = rows["Pine & Oak Studio"]
        self.assertTrue(str(ws.cell(row=prow, column=3).value).startswith("=IF(N(K"))
        # Amazon ana tabloda YOK
        self.assertNotIn("Amazon US - LazerCo", rows)


if __name__ == "__main__":
    unittest.main()
