# Nihai Excel raporu (ekibin şablon görünümü, openpyxl)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASLIKLAR = ["MAĞAZA", "CİRO", "Parça başı ciro", "VERGİ", "REKLAM",
             "KARGO MÜŞTERİ", "Kargo parça başı ort ödenen",
             "Kargo parça başı ort kalan", "Kargo", "P", "PARÇA ADEDİ",
             "İLAVE ÖDEME", "KALAN", "%REKLAM", "%VERGİ",
             "Parça başı kalan", "Upgrade"]

# Kolon harfleri (1 tabanlı): A MAĞAZA, B CİRO, C pb ciro, D VERGİ, E REKLAM,
# F KARGO MÜŞTERİ, G pb ödenen, H pb kalan, I Kargo, J P, K ADET, L İLAVE,
# M KALAN, N %REKLAM, O %VERGİ, P pb kalan, Q Upgrade
PARA_KOLONLARI = [2, 4, 5, 6, 9, 12, 13]
ONDALIK_KOLONLAR = [3, 7, 8, 16]
YUZDE_KOLONLARI = [14, 15]
SAYI_KOLONLARI = [11, 17]

TURUNCU = PatternFill("solid", fgColor="FF9900")
ACIK_GRI = PatternFill("solid", fgColor="F2F2F2")
BEYAZ_KALIN = Font(bold=True, color="FFFFFF", size=10)
KALIN = Font(bold=True)
INCE = Side(style="thin", color="BFBFBF")
KENARLIK = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)

PARA_FMT = "#,##0"
ONDALIK_FMT = "#,##0.00"
YUZDE_FMT = "0.0%"


def _formul_bolme(pay_hucre, adet_hucre):
    """ADET 0/boşsa boş bırakan bölme formülü (#DIV/0! asla üretilmez)."""
    return f'=IF(N({adet_hucre})=0,"",{pay_hucre}/{adet_hucre})'


def _stil(ws, satir, kalin=False, dolgu=None):
    for k in range(1, len(BASLIKLAR) + 1):
        h = ws.cell(row=satir, column=k)
        h.border = KENARLIK
        if kalin:
            h.font = KALIN
        if dolgu:
            h.fill = dolgu
        if k in PARA_KOLONLARI:
            h.number_format = PARA_FMT
        elif k in ONDALIK_KOLONLAR:
            h.number_format = ONDALIK_FMT
        elif k in YUZDE_KOLONLARI:
            h.number_format = YUZDE_FMT
        elif k in SAYI_KOLONLARI:
            h.number_format = "#,##0"


def rapor_uret(yol, ay_adi, yil, satirlar, amazon_satirlari=None,
               eslesmeyen_maliyet=None, urun_basliklari=None,
               siparis_icerik=None):
    """satirlar: [{magaza, ciro, vergi, reklam, kargo_musteri, kargo,
                   adet, ilave_odeme, upgrade, urunler?}]
    urun_basliklari verilirse rapor sağına ürün adet kolonları eklenir.
    siparis_icerik verilirse 2. sayfada çok-ürünlü siparişlerin içeriği
    listelenir (mağaza sahibinin Etsy'e manuel bakma ihtiyacını giderir)."""
    urun_basliklari = urun_basliklari or []
    son_fin = len(BASLIKLAR)  # son finansal kolon (Q = 17)
    tum_basliklar = list(BASLIKLAR) + list(urun_basliklari)

    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"

    # Başlık satırı (turuncu şablon) — finansal + ürün kolonları
    for k, ad in enumerate(tum_basliklar, start=1):
        h = ws.cell(row=1, column=k, value=ad)
        h.fill = TURUNCU
        h.font = BEYAZ_KALIN
        h.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="center" if k > 1 else "left")
        h.border = KENARLIK
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "B2"

    genislikler = [24, 11, 10, 10, 10, 12, 11, 11, 10, 4, 10, 10, 11, 9, 9, 11, 9]
    for k, g in enumerate(genislikler, start=1):
        ws.column_dimensions[get_column_letter(k)].width = g
    for k in range(son_fin + 1, len(tum_basliklar) + 1):
        ws.column_dimensions[get_column_letter(k)].width = 9

    # Ay bloğu başlığı
    r = 2
    ay_h = ws.cell(row=r, column=1, value=f"{ay_adi.upper()} {yil}")
    ay_h.font = Font(bold=True, size=12)
    _stil(ws, r, dolgu=ACIK_GRI)
    ay_h.font = Font(bold=True, size=12)
    r += 1

    ilk_veri = r
    for s in satirlar:
        ws.cell(row=r, column=1, value=s["magaza"])
        ws.cell(row=r, column=2, value=round(s.get("ciro") or 0, 2))
        ws.cell(row=r, column=4, value=round(s.get("vergi") or 0, 2))
        ws.cell(row=r, column=5, value=round(s.get("reklam") or 0, 2))
        ws.cell(row=r, column=6, value=round(s.get("kargo_musteri") or 0, 2))
        ws.cell(row=r, column=9, value=round(s.get("kargo") or 0, 2))
        # J ("P") kolonu kararla boş bırakılıyor
        ws.cell(row=r, column=11, value=s.get("adet") or 0)
        ws.cell(row=r, column=12, value=round(s.get("ilave_odeme") or 0, 2))
        ws.cell(row=r, column=17, value=s.get("upgrade") or 0)
        # Formüller (ekip hücreye tıklayıp kontrol edebilsin)
        ws.cell(row=r, column=3, value=_formul_bolme(f"B{r}", f"K{r}"))
        ws.cell(row=r, column=7, value=_formul_bolme(f"F{r}", f"K{r}"))
        ws.cell(row=r, column=8, value=_formul_bolme(f"(F{r}-I{r})", f"K{r}"))
        ws.cell(row=r, column=13, value=f"=B{r}-D{r}-E{r}-I{r}")
        ws.cell(row=r, column=14, value=f'=IF(N(B{r})=0,"",E{r}/B{r})')
        ws.cell(row=r, column=15, value=f'=IF(N(B{r})=0,"",D{r}/B{r})')
        ws.cell(row=r, column=16, value=_formul_bolme(f"M{r}", f"K{r}"))
        _stil(ws, r)
        # Ürün adet kolonları (form beyanından, sağ tarafta)
        urunler = s.get("urunler") or {}
        for j, ad in enumerate(urun_basliklari):
            c = ws.cell(row=r, column=son_fin + 1 + j, value=urunler.get(ad, 0))
            c.border = KENARLIK
            c.number_format = "#,##0"
        r += 1
    son_veri = r - 1

    # TOPLAM satırı
    if son_veri >= ilk_veri:
        ws.cell(row=r, column=1, value="TOPLAM")
        for k in [2, 4, 5, 6, 9, 11, 12, 13, 17]:
            harf = get_column_letter(k)
            ws.cell(row=r, column=k,
                    value=f"=SUM({harf}{ilk_veri}:{harf}{son_veri})")
        ws.cell(row=r, column=3, value=_formul_bolme(f"B{r}", f"K{r}"))
        ws.cell(row=r, column=7, value=_formul_bolme(f"F{r}", f"K{r}"))
        ws.cell(row=r, column=8, value=_formul_bolme(f"(F{r}-I{r})", f"K{r}"))
        # Yüzde kolonlarında toplam ciro üzerinden ağırlıklı değer
        ws.cell(row=r, column=14, value=f'=IF(N(B{r})=0,"",E{r}/B{r})')
        ws.cell(row=r, column=15, value=f'=IF(N(B{r})=0,"",D{r}/B{r})')
        ws.cell(row=r, column=16, value=_formul_bolme(f"M{r}", f"K{r}"))
        # Ürün kolonları toplamı
        for j in range(len(urun_basliklari)):
            harf = get_column_letter(son_fin + 1 + j)
            c = ws.cell(row=r, column=son_fin + 1 + j,
                        value=f"=SUM({harf}{ilk_veri}:{harf}{son_veri})")
            c.number_format = "#,##0"
        _stil(ws, r, kalin=True, dolgu=ACIK_GRI)
        # _stil yalnızca finansal kolonları gezdiği için ürün toplamlarına
        # kenarlık/kalınlığı ayrıca uygula
        for j in range(len(urun_basliklari)):
            c = ws.cell(row=r, column=son_fin + 1 + j)
            c.border = KENARLIK
            c.font = KALIN
            c.fill = ACIK_GRI
        r += 2

    # Rapor Dışı (Amazon) bölümü
    if amazon_satirlari:
        b = ws.cell(row=r, column=1, value="RAPOR DIŞI (AMAZON)")
        b.font = Font(bold=True, size=11)
        ws.cell(row=r, column=2, value="Sipariş/Gönderi")
        ws.cell(row=r, column=3, value="Kargo Harcaması")
        for k in (1, 2, 3):
            ws.cell(row=r, column=k).fill = ACIK_GRI
            ws.cell(row=r, column=k).border = KENARLIK
            ws.cell(row=r, column=k).font = KALIN
        r += 1
        for a in amazon_satirlari:
            ws.cell(row=r, column=1, value=a["magaza"]).border = KENARLIK
            ws.cell(row=r, column=2, value=a.get("siparis", 0)).border = KENARLIK
            m = ws.cell(row=r, column=3, value=round(a.get("kargo") or 0, 2))
            m.border = KENARLIK
            m.number_format = PARA_FMT
            r += 1
        r += 1

    if eslesmeyen_maliyet and eslesmeyen_maliyet.get("adet"):
        n = ws.cell(row=r, column=1, value=(
            f"Not: {eslesmeyen_maliyet['adet']} gönderinin mağazası "
            f"eşleştirilemedi (toplam {eslesmeyen_maliyet['maliyet']:.2f}$ kargo "
            "maliyeti rapora dahil edilmedi)."))
        n.font = Font(italic=True, color="990000")

    # 2. sayfa: Sipariş İçeriği (çok-ürünlü siparişler → gerçek ürünler)
    if siparis_icerik:
        ws2 = wb.create_sheet("Sipariş İçeriği")
        basliklar2 = ["Sipariş No", "Mağaza", "Adet", "Ürün", "SKU", "Kategori"]
        for k, ad in enumerate(basliklar2, start=1):
            h = ws2.cell(row=1, column=k, value=ad)
            h.fill = TURUNCU
            h.font = BEYAZ_KALIN
            h.border = KENARLIK
        ws2.freeze_panes = "A2"
        for k, g in zip([34, 22, 7, 60, 16, 18], range(1, 7)):
            ws2.column_dimensions[get_column_letter(g)].width = k
        r2 = 2
        for ono in sorted(siparis_icerik):
            g = siparis_icerik[ono]
            ilk = True
            for it in g["items"]:
                ws2.cell(row=r2, column=1,
                         value=ono if ilk else "").border = KENARLIK
                ws2.cell(row=r2, column=2,
                         value=g["store"] if ilk else "").border = KENARLIK
                ws2.cell(row=r2, column=3, value=it["adet"]).border = KENARLIK
                ws2.cell(row=r2, column=4, value=it["ad"]).border = KENARLIK
                ws2.cell(row=r2, column=5, value=it["sku"]).border = KENARLIK
                ws2.cell(row=r2, column=6, value=it["kategori"]).border = KENARLIK
                ilk = False
                r2 += 1

    wb.save(yol)
    return yol
